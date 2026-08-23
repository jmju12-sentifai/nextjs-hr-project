#!/usr/bin/env python3
"""
토스페이먼츠 결제경로 PPT 생성기.

가이드(토스페이먼츠_홈페이지_결제경로_제작_가이드_정기결제용(빌링).pdf)의
슬라이드 구성을 그대로 재현한다.

  표지 → ① 가맹점 정보 → ② 하단정보 → ③ 환불규정
       → ④ 로그인/회원가입 → ⑤ 상품선택/구매과정 → ⑥ 카드 결제경로

캡처 파일은 captures/ 에 아래 이름으로 넣는다. 여러 장이 필요한 단계는
숫자 접미사를 붙이면 순서대로 각각 슬라이드가 된다. (예: 05-purchase-1.png)

  02-footer.png          하단 사업자정보
  03-refund.png          환불정책 페이지
  04-login.png           로그인 / 회원가입
  05-purchase-*.png      요금 페이지 → 결제 페이지
  06-billing-*.png       카드 등록창(정기결제용 카드 입력창)

사용법:
    python3 scripts/build_payment_path_ppt.py
"""

from pathlib import Path
import re
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

from PIL import Image

# ── 가맹점 정보 (가이드 6p 표지 항목) ────────────────────────────────
MERCHANT = {
    "상호명": "케이프라임연구소",
    "사업자번호": "264-24-02200",
    "URL": "https://hrcoach.co.kr",
    "Test ID": "「테스트 계정 이메일」",
    "Test PW": "「테스트 계정 비밀번호」",
}

BLUE = RGBColor(0x4A, 0x72, 0xC4)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x80, 0x80, 0x80)
LINK = RGBColor(0x1B, 0x64, 0xD1)

ROOT = Path(__file__).resolve().parent.parent
CAPTURES = ROOT / "captures"
OUT = ROOT / "docs" / "케이프라임연구소_결제경로.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# (파일 접두사, 헤더 칩 문구, 부제)
STEPS = [
    ("02-footer", "② 하단 정보 캡처",
     "필수 구성 항목 : (1)상호명 / (2)대표자명 / (3)사업자등록번호 / "
     "(4)통신판매업신고번호 / (5)사업장주소 / (6)유선전화번호"),
    ("03-refund", "③ 환불규정 캡처",
     "정기결제 해지·환불 기준이 명시된 환불정책 페이지입니다."),
    ("04-login", "④ 로그인 / 회원가입 캡처",
     "비회원 구매가 불가하여 로그인 및 회원가입 경로를 함께 캡처하였습니다."),
    ("05-purchase", "⑤ 상품선택 / 구매과정 캡처",
     "상품의 명칭·상세설명·금액과 서비스 제공기간이 모두 확인됩니다."),
    ("06-billing", "⑥ 카드 결제경로 캡처",
     "빌링결제 정기결제용 카드 입력창까지 캡처하였습니다."),
]


def add_chip(slide, text):
    """상단 가운데 파란 라벨 칩."""
    w, h = Inches(3.2), Inches(0.42)
    box = slide.shapes.add_textbox((SLIDE_W - w) // 2, Inches(0.18), w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = BLUE
    box.line.fill.background()
    tf = box.text_frame
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_subtitle(slide, text):
    box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.72), SLIDE_W - Inches(1.2), Inches(0.5)
    )
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = INK


def add_image_fitted(slide, img_path, top=Inches(1.32)):
    """가로/세로 비율을 유지한 채 본문 영역에 맞춰 넣는다."""
    avail_w = SLIDE_W - Inches(1.0)
    avail_h = SLIDE_H - top - Inches(0.35)
    with Image.open(img_path) as im:
        iw, ih = im.size
    scale = min(avail_w / iw, avail_h / ih)
    w, h = Emu(int(iw * scale)), Emu(int(ih * scale))
    left = (SLIDE_W - w) // 2
    slide.shapes.add_picture(str(img_path), left, top, width=w, height=h)


def cover_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_chip(slide, "① 가맹점 정보 기재")

    box = slide.shapes.add_textbox(
        Inches(0.6), Inches(0.74), SLIDE_W - Inches(1.2), Inches(0.35)
    )
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "시작페이지에 가맹점 정보를 기재해요."
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = INK

    # 본문 카드
    card_l, card_t = Inches(1.1), Inches(1.5)
    card_w, card_h = SLIDE_W - Inches(2.2), Inches(5.4)
    card = slide.shapes.add_textbox(card_l, card_t, card_w, card_h)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    card.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    card.line.width = Pt(1)

    tf = slide.shapes.add_textbox(
        card_l + Inches(1.4), card_t + Inches(1.5),
        card_w - Inches(2.0), Inches(3.0)
    ).text_frame
    tf.word_wrap = True

    rows = [
        ("(1) 상호명", MERCHANT["상호명"], False),
        ("(2) 사업자번호", MERCHANT["사업자번호"], False),
        ("", "", False),
        ("(3) URL", MERCHANT["URL"], True),
        ("", "", False),
        ("(4) Test ID", MERCHANT["Test ID"], False),
        ("(5) Test PW", MERCHANT["Test PW"], False),
    ]
    first = True
    for label, value, is_link in rows:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(6)
        if not label:
            continue
        lr = p.add_run()
        lr.text = f"{label:<14}"
        lr.font.size = Pt(17)
        lr.font.bold = True
        lr.font.color.rgb = BLUE
        vr = p.add_run()
        vr.text = f" : {value}"
        vr.font.size = Pt(17)
        vr.font.bold = True
        vr.font.color.rgb = LINK if is_link else INK

    note = slide.shapes.add_textbox(
        Inches(0.6), SLIDE_H - Inches(0.62), SLIDE_W - Inches(1.2), Inches(0.4)
    ).text_frame
    p = note.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = (
        "상점아이디(MID) : bill_HRcoawdkw    ｜    "
        "결제수단 : 빌링결제(신용카드 정기결제)    ｜    "
        "서비스 제공기간 : 결제일로부터 1개월 (매월 자동갱신)"
    )
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED


def find_captures(prefix):
    """`05-purchase.png` 또는 `05-purchase-1.png` 형태를 순서대로 모은다."""
    exact = CAPTURES / f"{prefix}.png"
    if exact.exists():
        return [exact]

    def order(p):
        m = re.search(r"-(\d+)\.png$", p.name)
        return int(m.group(1)) if m else 0

    return sorted(CAPTURES.glob(f"{prefix}-*.png"), key=order)


def main():
    if not CAPTURES.exists():
        sys.exit(f"캡처 폴더가 없습니다: {CAPTURES}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    cover_slide(prs)

    missing = []
    for prefix, chip, subtitle in STEPS:
        shots = find_captures(prefix)
        if not shots:
            missing.append(prefix)
            continue
        for i, shot in enumerate(shots):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            label = chip if len(shots) == 1 else f"{chip}  ({i + 1}/{len(shots)})"
            add_chip(slide, label)
            add_subtitle(slide, subtitle)
            add_image_fitted(slide, shot)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)

    print(f"생성 완료: {OUT}")
    print(f"슬라이드 수: {len(prs.slides.__iter__.__self__._sldIdLst)}")
    if missing:
        print("\n⚠️ 캡처가 없어 건너뛴 단계:")
        for m in missing:
            print(f"   - {m}")


if __name__ == "__main__":
    main()
